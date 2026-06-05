import os
from jinja2 import Template
from pptx import Presentation
from pptx.util import Inches, Pt
from xhtml2pdf import pisa

# Sample SEO Data Dictionary
seo_data = {
    "project_name": "ForgePractice SEO Audit",
    "date": "2026-06-05",
    "metrics": {
        "organic_traffic": "12,500",
        "avg_position": "14.2",
        "backlinks": "450",
        "domain_authority": "32"
    },
    "top_keywords": [
        {"keyword": "python automation", "position": 5, "volume": 1200},
        {"keyword": "seo reporting tool", "position": 12, "volume": 800},
        {"keyword": "weasyprint pdf", "position": 3, "volume": 500},
        {"keyword": "pptx generation python", "position": 8, "volume": 300},
    ],
    "recommendations": [
        "Optimize meta titles for top 5 pages",
        "Increase internal linking to high-conversion pages",
        "Fix 404 errors identified in the crawl",
        "Improve page load speed for mobile devices"
    ]
}

def generate_html(data):
    print("Generating HTML report...")
    html_template = """
    <!DOCTYPE html>
    <html>
    <head>
        <title>{{ project_name }}</title>
        <style>
            body { font-family: Arial, sans-serif; margin: 40px; color: #333; }
            h1 { color: #2c3e50; }
            .metric-box { margin-bottom: 30px; }
            .metric { background: #f4f7f6; padding: 15px; border-radius: 8px; display: inline-block; width: 20%; text-align: center; border: 1px solid #ddd; margin-right: 10px; }
            .metric-value { font-size: 24px; font-weight: bold; color: #e67e22; }
            table { width: 100%; border-collapse: collapse; margin-bottom: 30px; }
            th, td { border: 1px solid #ddd; padding: 12px; text-align: left; }
            th { background-color: #f2f2f2; }
            .recommendations { background: #fffbe6; padding: 20px; border-left: 5px solid #f1c40f; }
        </style>
    </head>
    <body>
        <h1>{{ project_name }} - SEO Report</h1>
        <p>Date: {{ date }}</p>

        <div class="metric-box">
            {% for key, value in metrics.items() %}
            <div class="metric">
                <div style="text-transform: capitalize;">{{ key.replace('_', ' ') }}</div>
                <div class="metric-value">{{ value }}</div>
            </div>
            {% endfor %}
        </div>

        <h2>Top Keywords</h2>
        <table>
            <thead>
                <tr>
                    <th>Keyword</th>
                    <th>Position</th>
                    <th>Volume</th>
                </tr>
            </thead>
            <tbody>
                {% for item in top_keywords %}
                <tr>
                    <td>{{ item.keyword }}</td>
                    <td>{{ item.position }}</td>
                    <td>{{ item.volume }}</td>
                </tr>
                {% endfor %}
            </tbody>
        </table>

        <h2>Recommendations</h2>
        <div class="recommendations">
            <ul>
                {% for rec in recommendations %}
                <li>{{ rec }}</li>
                {% endfor %}
            </ul>
        </div>
    </body>
    </html>
    """
    template = Template(html_template)
    rendered_html = template.render(data)

    with open("report.html", "w", encoding="utf-8") as f:
        f.write(rendered_html)
    return "report.html"

def generate_pdf(html_file):
    print("Generating PDF report (using xhtml2pdf)...")
    # Note: xhtml2pdf is used as a fallback for weasyprint on Windows
    # due to the requirement of system-level GTK libraries.
    with open(html_file, "rb") as f:
        html_content = f.read()

    with open("report.pdf", "wb") as f:
        pisa.CreatePDF(html_content, f)
    return "report.pdf"

def generate_pptx(data):
    print("Generating PPTX report...")
    prs = Presentation()

    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = data["project_name"]
    subtitle.text = f"SEO Analysis Report\nDate: {data['date']}"

    # Metrics Slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    title_shape.text = "Key Metrics"

    body_shape = shapes.placeholders[1]
    tf = body_shape.text_frame
    for key, value in data["metrics"].items():
        p = tf.add_paragraph()
        p.text = f"{key.replace('_', ' ').title()}: {value}"
        p.level = 0

    # Keywords Slide
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Top Keywords"
    tf = slide.placeholders[1].text_frame
    for item in data["top_keywords"]:
        p = tf.add_paragraph()
        p.text = f"{item['keyword']} - Pos: {item['position']} (Vol: {item['volume']})"
        p.level = 0

    # Recommendations Slide
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Recommendations"
    tf = slide.placeholders[1].text_frame
    for rec in data["recommendations"]:
        p = tf.add_paragraph()
        p.text = rec
        p.level = 0

    prs.save("report.pptx")
    return "report.pptx"

if __name__ == "__main__":
    try:
        html_path = generate_html(seo_data)
        pdf_path = generate_pdf(html_path)
        pptx_path = generate_pptx(seo_data)
        print(f"Successfully generated:\n- {html_path}\n- {pdf_path}\n- {pptx_path}")
    except Exception as e:
        print(f"Error during generation: {e}")
        import traceback
        traceback.print_exc()
